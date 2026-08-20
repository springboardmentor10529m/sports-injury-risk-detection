import io
import os
from datetime import datetime
from typing import Dict, Any, List

# ReportLab imports for PDF generation (optional dependency)
try:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, HRFlowable
    )
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

# OpenPyXL imports for Excel generation (optional dependency)
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False



def generate_pdf_report(data: Dict[str, Any]) -> bytes:
    """
    Generates a professional PDF injury risk assessment report using ReportLab.
    """
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError(
            "reportlab is not installed. Run: pip install reportlab"
        )
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    
    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor("#0f172a"),
        alignment=1  # Centered
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#64748b"),
        alignment=1
    )

    h2_style = ParagraphStyle(
        'Heading2Custom',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#0d9488"),
        spaceBefore=10,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'BodyCustom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#334155")
    )

    disclaimer_style = ParagraphStyle(
        'DisclaimerCustom',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#ef4444")
    )

    elements = []

    # Title Banner
    elements.append(Paragraph("AI SPORTS INJURY RISK ASSESSMENT REPORT", title_style))
    elements.append(Spacer(1, 4))
    elements.append(Paragraph("Biomechanical Analysis & Predictive Injury Risk Screening System", subtitle_style))
    elements.append(Spacer(1, 10))
    elements.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor("#0d9488"), spaceAfter=15))

    # Athlete Information Header Table
    athlete_info = data.get("athlete", {})
    video_info = data.get("video", {})
    created_at = data.get("created_at", datetime.now().strftime("%Y-%m-%d %H:%M"))

    info_data = [
        [
            Paragraph(f"<b>Athlete Name:</b> {athlete_info.get('name', 'N/A')}", body_style),
            Paragraph(f"<b>Sport & Position:</b> {athlete_info.get('sport', 'General')} ({athlete_info.get('position', 'N/A')})", body_style)
        ],
        [
            Paragraph(f"<b>Age / Height / Weight:</b> {athlete_info.get('age', 'N/A')} yrs | {athlete_info.get('height', 'N/A')} cm | {athlete_info.get('weight', 'N/A')} kg", body_style),
            Paragraph(f"<b>Assessment Date:</b> {created_at}", body_style)
        ],
        [
            Paragraph(f"<b>Activity Assessed:</b> {video_info.get('activity', 'Squatting')}", body_style),
            Paragraph(f"<b>Training Load:</b> {athlete_info.get('training_load', 0.0)} hrs/week", body_style)
        ]
    ]

    info_table = Table(info_data, colWidths=[270, 270])
    info_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#cbd5e1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 15))

    # Risk Summary Header Card
    risk_data = data.get("risk", {})
    overall_score = risk_data.get("risk_score", 0.0)
    risk_category = risk_data.get("risk_category", "Low")
    
    cat_color = colors.HexColor("#10b981") if risk_category == "Low" else (
        colors.HexColor("#f59e0b") if risk_category == "Moderate" else (
            colors.HexColor("#ef4444") if risk_category == "High" else colors.HexColor("#991b1b")
        )
    )

    summary_data = [
        [
            Paragraph("<b>Overall Injury Risk Score</b>", ParagraphStyle('W1', parent=body_style, textColor=colors.white, fontName="Helvetica-Bold", fontSize=11)),
            Paragraph("<b>Risk Category Classification</b>", ParagraphStyle('W2', parent=body_style, textColor=colors.white, fontName="Helvetica-Bold", fontSize=11))
        ],
        [
            Paragraph(f"<font size=18><b>{overall_score}%</b></font>", ParagraphStyle('Score', parent=body_style, textColor=colors.white, alignment=1)),
            Paragraph(f"<font size=16><b>{risk_category.upper()} RISK</b></font>", ParagraphStyle('Cat', parent=body_style, textColor=colors.white, alignment=1))
        ]
    ]

    summary_table = Table(summary_data, colWidths=[270, 270])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), cat_color),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 15))

    # Biomechanical Metrics Section
    elements.append(Paragraph("Biomechanical Measurements & Movement Quality", h2_style))
    assessment = data.get("assessment", {})

    bio_rows = [
        ["Metric", "Measured Value", "Target / Reference", "Status"],
        ["Knee Valgus Ratio", f"{assessment.get('knee_valgus', 'N/A')}", ">= 0.85", "Normal" if (assessment.get('knee_valgus', 1.0) or 1.0) >= 0.85 else "Collapsed"],
        ["Hip Tilt (Stability)", f"{assessment.get('hip_stability', 'N/A')} deg", "<= 5.0 deg", "Stable" if (assessment.get('hip_stability', 0.0) or 0) <= 5.0 else "Lateral Drop"],
        ["Trunk Forward Lean", f"{assessment.get('trunk_lean', 'N/A')} deg", "<= 20.0 deg", "Good Control" if (assessment.get('trunk_lean', 0.0) or 0) <= 20.0 else "Excessive Lean"],
        ["Limb Symmetry Score", f"{assessment.get('symmetry_score', 'N/A')}%", ">= 90.0%", "Symmetric" if (assessment.get('symmetry_score', 100) or 100) >= 90.0 else "Asymmetric"],
        ["Fatigue Score", f"{assessment.get('fatigue_score', 'N/A')} / 10", "<= 4.0", "Low Fatigue" if (assessment.get('fatigue_score', 0) or 0) <= 4.0 else "Elevated Fatigue"],
        ["Movement Quality Score", f"{assessment.get('movement_quality', 'N/A')}%", ">= 75.0%", "Optimal" if (assessment.get('movement_quality', 100) or 100) >= 75.0 else "Sub-optimal"]
    ]

    bio_table = Table(bio_rows, colWidths=[150, 120, 130, 140])
    bio_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0f766e")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(bio_table)
    elements.append(Spacer(1, 15))

    # Predicted Injury Risks Breakdown
    elements.append(Paragraph("Predicted Specific Injury Probabilities", h2_style))
    
    risk_rows = [
        ["Injury Classification Risk", "Probability Score", "Risk Assessment Level"],
        ["ACL Injury Risk", f"{risk_data.get('acl_risk', 0.0)}%", "Elevated" if risk_data.get('acl_risk', 0.0) > 50 else "Low"],
        ["Hamstring Strain Risk", f"{risk_data.get('hamstring_risk', 0.0)}%", "Elevated" if risk_data.get('hamstring_risk', 0.0) > 50 else "Low"],
        ["Ankle Sprain Risk", f"{risk_data.get('ankle_risk', 0.0)}%", "Elevated" if risk_data.get('ankle_risk', 0.0) > 50 else "Low"],
        ["Shoulder Injury Risk", f"{risk_data.get('shoulder_risk', 0.0)}%", "Elevated" if risk_data.get('shoulder_risk', 0.0) > 50 else "Low"],
        ["Lower Back Strain Risk", f"{risk_data.get('lower_back_risk', 0.0)}%", "Elevated" if risk_data.get('lower_back_risk', 0.0) > 50 else "Low"],
        ["Overuse Injury Risk", f"{risk_data.get('overuse_risk', 0.0)}%", "Elevated" if risk_data.get('overuse_risk', 0.0) > 50 else "Low"]
    ]

    risk_table = Table(risk_rows, colWidths=[200, 160, 180])
    risk_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(risk_table)
    elements.append(Spacer(1, 15))

    # Corrective Recommendations Section
    elements.append(Paragraph("Personalized AI Corrective Recommendations", h2_style))
    recs = data.get("recommendations", {})

    rec_content = []
    if recs.get("exercise"):
        rec_content.append(f"<b>Prescribed Exercises:</b><br/>{recs['exercise'].replace(chr(10), '<br/>')}")
    if recs.get("mobility"):
        rec_content.append(f"<b>Mobility Protocol:</b><br/>{recs['mobility'].replace(chr(10), '<br/>')}")
    if recs.get("strengthening"):
        rec_content.append(f"<b>Strengthening Targets:</b><br/>{recs['strengthening'].replace(chr(10), '<br/>')}")
    if recs.get("recovery"):
        rec_content.append(f"<b>Recovery Guidelines:</b><br/>{recs['recovery'].replace(chr(10), '<br/>')}")
    if recs.get("training_modification"):
        rec_content.append(f"<b>Training Modifications:</b><br/>{recs['training_modification']}")

    for rec_text in rec_content:
        elements.append(Paragraph(rec_text, body_style))
        elements.append(Spacer(1, 6))

    elements.append(Spacer(1, 15))
    elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#cbd5e1"), spaceAfter=10))
    
    # Medical Disclaimer Footer
    disclaimer_text = (
        "<b>MEDICAL DISCLAIMER:</b> This platform uses AI-based computer vision for movement screening and biomechanical risk detection. "
        "The generated scores and recommendations are screening indicators and DO NOT constitute a medical diagnosis, clinical prognosis, "
        "or treatment prescription. Always consult a qualified physiotherapist, sports scientist, or medical physician before beginning "
        "rehabilitation or altering athletic training programs."
    )
    elements.append(Paragraph(disclaimer_text, disclaimer_style))

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()


def generate_excel_report(data: Dict[str, Any]) -> bytes:
    """
    Generates a multi-tab structured Excel workbook for injury risk data using OpenPyXL.
    """
    if not OPENPYXL_AVAILABLE:
        raise RuntimeError(
            "openpyxl is not installed. Run: pip install openpyxl"
        )
    wb = openpyxl.Workbook()
    
    header_fill = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    athlete_info = data.get("athlete", {})
    risk_data = data.get("risk", {})
    assessment = data.get("assessment", {})
    recs = data.get("recommendations", {})
    anomalies = data.get("anomalies", [])

    # Sheet 1: Assessment Summary
    ws1 = wb.active
    ws1.title = "Assessment Summary"

    ws1.append(["SPORTS INJURY RISK ASSESSMENT SUMMARY"])
    ws1.merge_cells("A1:D1")
    ws1["A1"].font = Font(name="Calibri", size=14, bold=True, color="0D9488")
    ws1.append([])

    summary_rows = [
        ["Athlete Name", athlete_info.get("name", "N/A"), "Sport", athlete_info.get("sport", "General")],
        ["Position", athlete_info.get("position", "N/A"), "Assessment Date", data.get("created_at", str(datetime.now()))],
        ["Age", athlete_info.get("age", "N/A"), "Height (cm)", athlete_info.get("height", "N/A")],
        ["Weight (kg)", athlete_info.get("weight", "N/A"), "Training Load (hrs/wk)", athlete_info.get("training_load", 0.0)],
        [],
        ["OVERALL INJURY RISK SCORE", f"{risk_data.get('risk_score', 0.0)}%", "RISK CATEGORY", risk_data.get("risk_category", "Low")],
        ["MOVEMENT QUALITY SCORE", f"{assessment.get('movement_quality', 100)}%", "SYMMETRY SCORE", f"{assessment.get('symmetry_score', 100)}%"]
    ]

    for row in summary_rows:
        ws1.append(row)

    # Style Sheet 1
    for r in range(3, 10):
        for c in range(1, 5):
            cell = ws1.cell(row=r, column=c)
            if c in (1, 3):
                cell.font = bold_font
            else:
                cell.font = regular_font
            cell.border = thin_border

    # Sheet 2: Biomechanics & Risks
    ws2 = wb.create_sheet(title="Biomechanics & Risks")
    
    ws2.append(["Biomechanical Metric", "Measured Value", "Reference Range", "Status"])
    for col in range(1, 5):
        cell = ws2.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font

    bio_data = [
        ["Knee Valgus Ratio", assessment.get("knee_valgus", "N/A"), ">= 0.85", "Normal" if (assessment.get('knee_valgus', 1.0) or 1.0) >= 0.85 else "Collapsed"],
        ["Hip Tilt (Stability)", f"{assessment.get('hip_stability', 'N/A')} deg", "<= 5.0 deg", "Stable" if (assessment.get('hip_stability', 0.0) or 0) <= 5.0 else "Lateral Drop"],
        ["Trunk Lean", f"{assessment.get('trunk_lean', 'N/A')} deg", "<= 20.0 deg", "Good Control" if (assessment.get('trunk_lean', 0.0) or 0) <= 20.0 else "Excessive Lean"],
        ["Limb Symmetry Score", f"{assessment.get('symmetry_score', 'N/A')}%", ">= 90.0%", "Symmetric" if (assessment.get('symmetry_score', 100) or 100) >= 90.0 else "Asymmetric"],
        ["Fatigue Score", f"{assessment.get('fatigue_score', 'N/A')} / 10", "<= 4.0", "Low Fatigue" if (assessment.get('fatigue_score', 0) or 0) <= 4.0 else "Elevated"]
    ]

    for row in bio_data:
        ws2.append(row)

    ws2.append([])
    ws2.append(["Specific Injury Risk", "Probability Score (%)", "Risk Level"])
    header_row_idx = ws2.max_row
    for col in range(1, 4):
        cell = ws2.cell(row=header_row_idx, column=col)
        cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        cell.font = header_font

    risk_breakdown = [
        ["ACL Injury Risk", risk_data.get("acl_risk", 0.0), "High" if risk_data.get("acl_risk", 0.0) > 50 else "Low"],
        ["Hamstring Strain Risk", risk_data.get("hamstring_risk", 0.0), "High" if risk_data.get("hamstring_risk", 0.0) > 50 else "Low"],
        ["Ankle Sprain Risk", risk_data.get("ankle_risk", 0.0), "High" if risk_data.get("ankle_risk", 0.0) > 50 else "Low"],
        ["Shoulder Injury Risk", risk_data.get("shoulder_risk", 0.0), "High" if risk_data.get("shoulder_risk", 0.0) > 50 else "Low"],
        ["Lower Back Strain Risk", risk_data.get("lower_back_risk", 0.0), "High" if risk_data.get("lower_back_risk", 0.0) > 50 else "Low"],
        ["Overuse Injury Risk", risk_data.get("overuse_risk", 0.0), "High" if risk_data.get("overuse_risk", 0.0) > 50 else "Low"]
    ]

    for row in risk_breakdown:
        ws2.append(row)

    # Sheet 3: Movement Anomalies
    ws3 = wb.create_sheet(title="Movement Anomalies")
    ws3.append(["Start (s)", "End (s)", "Issue Type", "Severity", "Confidence", "Affected Joints", "Description"])
    for col in range(1, 8):
        cell = ws3.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font

    if anomalies:
        for a in anomalies:
            ws3.append([
                a.get("timestamp_start", 0.0),
                a.get("timestamp_end", 0.0),
                a.get("issue_type", "Deviation"),
                a.get("severity", "Low"),
                a.get("confidence", 0.85),
                a.get("affected_joints", "Joints"),
                a.get("description", "")
            ])
    else:
        ws3.append([0.0, 0.0, "No Severe Anomalies", "Low", 1.0, "None", "Movement pattern falls within normal physiological limits."])

    # Sheet 4: Corrective Plan
    ws4 = wb.create_sheet(title="Corrective Recommendations")
    ws4.append(["Category", "Details & Protocol"])
    for col in range(1, 3):
        cell = ws4.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font

    plan_rows = [
        ["Prescribed Exercises", recs.get("exercise", "Maintain current regime.")],
        ["Mobility Suggestions", recs.get("mobility", "Standard dynamic warm-up.")],
        ["Strengthening Targets", recs.get("strengthening", "Bodyweight core stability.")],
        ["Recovery Plan", recs.get("recovery", "Maintain sleeping & hydration routines.")],
        ["Training Modifications", recs.get("training_modification", "No load reductions required.")]
    ]

    for row in plan_rows:
        ws4.append(row)

    # Auto-adjust column widths for all worksheets
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                max_len = max(max_len, len(val))
            sheet.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
